"""把项目知识文档转换为可索引的文本块。"""

from __future__ import annotations

import hashlib
import unicodedata
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable


SUPPORTED_EXTENSIONS = (".md", ".txt", ".pdf", ".docx", ".pptx")


@dataclass(frozen=True)
class KnowledgeDocument:
    content: str
    source: str
    source_path: str
    page: int | None = None
    kind: str = "text"


@dataclass(frozen=True)
class KnowledgeChunk:
    chunk_id: str
    content: str
    metadata: dict[str, Any]


def normalize_text(text: str) -> str:
    text = unicodedata.normalize("NFKC", text.replace("\x00", ""))
    lines = [" ".join(line.split()) for line in text.splitlines()]
    compact: list[str] = []
    previous_blank = False
    for line in lines:
        if line:
            compact.append(line)
            previous_blank = False
        elif not previous_blank and compact:
            compact.append("")
            previous_blank = True
    return "\n".join(compact).strip()


def discover_knowledge_files(knowledge_dir: Path, include_files: Iterable[Path]) -> list[Path]:
    files: dict[str, Path] = {}
    if knowledge_dir.is_dir():
        for path in knowledge_dir.rglob("*"):
            if path.is_file() and path.suffix.lower() in SUPPORTED_EXTENSIONS:
                files[str(path.resolve()).casefold()] = path.resolve()
    for path in include_files:
        resolved = path.resolve()
        if resolved.is_file() and resolved.suffix.lower() in SUPPORTED_EXTENSIONS:
            files[str(resolved).casefold()] = resolved
    return sorted(files.values(), key=lambda item: str(item).casefold())


def _source_name(path: Path, project_root: Path) -> str:
    try:
        return path.relative_to(project_root).as_posix()
    except ValueError:
        return path.name


def _load_plain(path: Path, project_root: Path) -> list[KnowledgeDocument]:
    content = path.read_text(encoding="utf-8-sig", errors="replace")
    return [KnowledgeDocument(normalize_text(content), _source_name(path, project_root), str(path), kind=path.suffix[1:])]


def _load_pdf(path: Path, project_root: Path) -> list[KnowledgeDocument]:
    try:
        from pypdf import PdfReader
    except ImportError as error:
        raise RuntimeError("读取PDF需要安装 pypdf") from error
    source = _source_name(path, project_root)
    reader = PdfReader(str(path))
    return [
        KnowledgeDocument(normalize_text(page.extract_text() or ""), source, str(path), number, "pdf")
        for number, page in enumerate(reader.pages, start=1)
        if normalize_text(page.extract_text() or "")
    ]


def _load_docx(path: Path, project_root: Path) -> list[KnowledgeDocument]:
    try:
        from docx import Document
    except ImportError as error:
        raise RuntimeError("读取DOCX需要安装 python-docx") from error
    document = Document(str(path))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        parts.extend(" | ".join(cell.text for cell in row.cells) for row in table.rows)
    return [KnowledgeDocument(normalize_text("\n".join(parts)), _source_name(path, project_root), str(path), kind="docx")]


def _pptx_shape_text(shape) -> list[str]:
    parts: list[str] = []
    if getattr(shape, "shape_type", None) == 6 and hasattr(shape, "shapes"):
        for child in shape.shapes:
            parts.extend(_pptx_shape_text(child))
    if getattr(shape, "has_text_frame", False):
        value = normalize_text(shape.text)
        if value:
            parts.append(value)
    if getattr(shape, "has_table", False):
        for row in shape.table.rows:
            value = normalize_text(" | ".join(cell.text for cell in row.cells))
            if value:
                parts.append(value)
    return parts


def _load_pptx(path: Path, project_root: Path) -> list[KnowledgeDocument]:
    try:
        from pptx import Presentation
    except ImportError as error:
        raise RuntimeError("读取PPTX需要安装 python-pptx") from error
    source = _source_name(path, project_root)
    presentation = Presentation(str(path))
    documents: list[KnowledgeDocument] = []
    for number, slide in enumerate(presentation.slides, start=1):
        blocks: list[tuple[int, int, str]] = []
        for shape in slide.shapes:
            value = normalize_text("\n".join(_pptx_shape_text(shape)))
            if value:
                blocks.append((int(shape.top), int(shape.left), value))
        content = normalize_text("\n".join(item[2] for item in sorted(blocks)))
        if content:
            documents.append(KnowledgeDocument(content, source, str(path), number, "pptx"))
    return documents


def load_knowledge_file(path: Path, project_root: Path) -> list[KnowledgeDocument]:
    loaders = {
        ".md": _load_plain,
        ".txt": _load_plain,
        ".pdf": _load_pdf,
        ".docx": _load_docx,
        ".pptx": _load_pptx,
    }
    try:
        loader = loaders[path.suffix.lower()]
    except KeyError as error:
        raise ValueError(f"不支持的知识文档: {path}") from error
    return [document for document in loader(path, project_root) if document.content]


def _split_text(text: str, chunk_size: int, overlap: int) -> list[str]:
    if len(text) <= chunk_size:
        return [text]
    chunks: list[str] = []
    start = 0
    while start < len(text):
        proposed = min(start + chunk_size, len(text))
        end = proposed
        if proposed < len(text):
            minimum = start + max(chunk_size // 2, 1)
            candidates = [text.rfind(separator, minimum, proposed) for separator in ("\n\n", "\n", "。", "；")]
            boundary = max(candidates)
            if boundary >= minimum:
                end = boundary + 1
        value = text[start:end].strip()
        if value:
            chunks.append(value)
        if end >= len(text):
            break
        start = max(end - overlap, start + 1)
    return chunks


def chunk_documents(
    documents: Iterable[KnowledgeDocument],
    *,
    chunk_size: int,
    overlap: int,
) -> list[KnowledgeChunk]:
    chunks: list[KnowledgeChunk] = []
    for document in documents:
        location = f"第{document.page}页" if document.page else ""
        prefix = f"来源：{document.source}" + (f"\n位置：{location}" if location else "")
        for index, text in enumerate(_split_text(document.content, chunk_size, overlap), start=1):
            content = f"{prefix}\n\n{text}"
            digest = hashlib.sha256(
                f"{document.source}|{document.page}|{index}|{content}".encode("utf-8")
            ).hexdigest()
            chunks.append(
                KnowledgeChunk(
                    chunk_id=digest,
                    content=content,
                    metadata={
                        "source": document.source,
                        "source_path": document.source_path,
                        "page": document.page or 0,
                        "kind": document.kind,
                        "chunk": index,
                        "citation": f"[{document.source}{f'，第{document.page}页' if document.page else ''}]",
                    },
                )
            )
    return chunks
